"""Extract text, numbers and key terms per slide from PDF or PPTX."""

import re
from collections import Counter
from pathlib import Path
from typing import TypedDict

from .tokens import tokens

MIN_WORDS = 5  # fewer words than this: image-only slide, left out of the score
MAX_KEY_TERMS = 8
MIN_TERM_LEN = 5

NUMBER_RE = re.compile(
    r"\d+(?:[.,]\d+)?\s*(?:%|хувь|сая|тэрбум|мянга|₮|\$|k|m)?",
    re.IGNORECASE,
)

STOPWORDS = {
    "болон", "байна", "байгаа", "бөгөөд", "зэрэг", "гэсэн", "гэдэг", "тухай", "хамгийн",
    "бидний", "манай", "таны", "эдгээр", "энэхүү", "which", "there", "about", "their", "these",
}


class SlideText(TypedDict):
    number: int
    text: str
    notes: str
    image_only: bool
    numbers: list[str]
    key_terms: list[str]


class SlidesError(Exception):
    pass


def extract(path: Path) -> list[SlideText]:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            raw = _pdf(path)
        elif suffix == ".pptx":
            raw = _pptx(path)
        else:
            raise SlidesError("Slides must be a PDF or PPTX file.")
    except SlidesError:
        raise
    except Exception as e:
        raise SlidesError("Could not read the slides file. Is it a valid PDF or PPTX?") from e

    return [_describe(i + 1, text, notes) for i, (text, notes) in enumerate(raw)]


def _pdf(path: Path) -> list[tuple[str, str]]:
    import fitz  # PyMuPDF

    with fitz.open(path) as doc:
        return [(_clean(page.get_text()), "") for page in doc]


def _pptx(path: Path) -> list[tuple[str, str]]:
    from pptx import Presentation

    out = []
    for slide in Presentation(str(path)).slides:
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.extend(cell.text for cell in row.cells)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
        out.append((_clean("\n".join(parts)), _clean(notes)))
    return out


def _clean(text: str) -> str:
    return re.sub(r"[ \t]+", " ", re.sub(r"\n{2,}", "\n", text)).strip()


def _describe(number: int, text: str, notes: str) -> SlideText:
    words = tokens(text)
    numbers = []
    for m in NUMBER_RE.finditer(text):
        value = m.group(0).strip()
        has_unit = not value.replace(",", "").replace(".", "").isdigit()
        if has_unit or len(re.sub(r"\D", "", value)) >= 2:
            numbers.append(value)

    counts = Counter(w for w in words if len(w) >= MIN_TERM_LEN and not w.isdigit() and w not in STOPWORDS)
    key_terms = [w for w, _ in counts.most_common(MAX_KEY_TERMS)]

    return {
        "number": number,
        "text": text,
        "notes": notes,
        "image_only": len(words) < MIN_WORDS,
        "numbers": list(dict.fromkeys(numbers)),
        "key_terms": key_terms,
    }
